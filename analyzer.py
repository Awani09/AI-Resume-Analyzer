"""Resume text extraction and explainable job matching."""
from pathlib import Path
import re

import fitz  # PyMuPDF
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Edit or extend this vocabulary to fit the roles you want to analyze.
SKILLS = [
    "python", "java", "c#", ".net", "asp.net", "javascript", "typescript",
    "html", "css", "bootstrap", "react", "angular", "node.js", "flask",
    "django", "sql", "mysql", "postgresql", "sqlite", "mongodb",
    "machine learning", "deep learning", "natural language processing",
    "nlp", "data analysis", "pandas", "numpy", "scikit-learn", "tensorflow",
    "pytorch", "excel", "power bi", "tableau", "git", "github", "rest api",
    "api", "aws", "azure", "docker", "linux", "oops", "object-oriented programming",
    "data structures", "algorithms", "testing", "debugging", "agile",
    "communication", "problem solving", "teamwork"
]


def extract_text(file_path: str, extension: str) -> str:
    """Extract text from a PDF or DOCX file."""
    extension = extension.lower()
    if extension == ".pdf":
        chunks = []
        with fitz.open(file_path) as pdf:
            for page in pdf:
                chunks.append(page.get_text("text"))
        text = "\n".join(chunks)
    elif extension == ".docx":
        document = Document(file_path)
        chunks = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(chunks)
    else:
        raise ValueError("Unsupported file type. Upload a PDF or DOCX file.")

    cleaned = re.sub(r"[ \t]+", " ", text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    if len(cleaned) < 30:
        raise ValueError(
            "Very little text could be extracted. If this is a scanned PDF, "
            "try a text-based PDF or DOCX file. OCR is not included yet."
        )
    return cleaned


def _find_skills(text: str) -> list[str]:
    """Find known skills using case-insensitive whole-term matching."""
    lower = text.lower()
    found = []
    for skill in SKILLS:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    return sorted(set(found), key=str.lower)


def analyze_resume(resume_text: str, job_description: str) -> dict:
    """Return explainable score components, skill gaps, and suggestions."""
    if not resume_text.strip() or not job_description.strip():
        raise ValueError("Resume text and job description are both required.")

    resume_skills = set(_find_skills(resume_text))
    job_skills = set(_find_skills(job_description))
    matched = sorted(resume_skills & job_skills, key=str.lower)
    missing = sorted(job_skills - resume_skills, key=str.lower)

    try:
        matrix = TfidfVectorizer(stop_words="english", ngram_range=(1, 2)).fit_transform(
            [resume_text, job_description]
        )
        similarity = float(cosine_similarity(matrix[0:1], matrix[1:2])[0][0])
    except ValueError:
        similarity = 0.0

    if job_skills:
        skill_coverage = len(matched) / len(job_skills)
        score = round((skill_coverage * 70) + (similarity * 30))
        score_method = "70% detected job-skill coverage + 30% text similarity"
    else:
        skill_coverage = None
        score = round(similarity * 100)
        score_method = "Text similarity only; no recognized skills were detected in the job description"

    suggestions = []
    if missing:
        suggestions.append(
            "Review the missing skills and add any that you genuinely have, supported by a project, course, or work example."
        )
    if similarity < 0.20:
        suggestions.append(
            "Tailor the summary and project descriptions to the responsibilities in the job description, using truthful and relevant wording."
        )
    if not re.search(r"\b(project|projects|experience|internship|developed|built|created)\b", resume_text, re.I):
        suggestions.append(
            "Consider adding a Projects or Experience section with specific contributions and outcomes, if applicable."
        )
    if not suggestions:
        suggestions.append(
            "Your resume contains several relevant terms. Improve it further by showing evidence of those skills through specific projects and outcomes."
        )

    return {
        "score": max(0, min(100, score)),
        "similarity_percent": round(similarity * 100),
        "skill_coverage_percent": round(skill_coverage * 100) if skill_coverage is not None else None,
        "matched_skills": matched,
        "missing_skills": missing,
        "resume_skills": sorted(resume_skills, key=str.lower),
        "job_skills": sorted(job_skills, key=str.lower),
        "suggestions": suggestions,
        "score_method": score_method,
    }
